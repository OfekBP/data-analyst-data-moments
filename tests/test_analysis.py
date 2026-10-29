"""Proof that the numbers in the answers are right, and that the code adapts when the numbers change.

Three kinds of test:
  INVARIANTS  must hold for ANY data with this layout (SQL == pandas, parts add up to the whole, ...).
  ADAPTIVITY  change the figures on purpose and check the conclusions follow.
  REFERENCE   pin the headline figures of the delivered dataset (the ones quoted in docs/answers.md).
              If you change the data or the logic these fail ON PURPOSE - the message says so. It is a
              reminder to re-read docs/answers.md, not a sign that the code is broken.
"""
import pytest

from dm_analyst import config, questions, run_analysis, sql_check
from dm_analyst.load import load_raw


@pytest.fixture(scope="module")
def res():
    return run_analysis.compute_all()


@pytest.fixture(scope="module")
def sql(res):
    return sql_check.run_all(res["df"])


def pin(name, actual, expected, abs_tol):
    """Reference-figure check with an explanation instead of a bare assertion error."""
    assert actual == pytest.approx(expected, abs=abs_tol), (
        f"\n  REFERENCE FIGURE CHANGED: {name} is now {actual!r}, the delivered analysis had {expected!r}."
        "\n  If you changed the data or the code on purpose this is expected: run"
        "\n  `python -m dm_analyst.run_analysis` to see everything that moved, then update docs/answers.md"
        "\n  and the reference value in this test. If you did not, something broke."
    )


# ============================================================== INVARIANTS
def test_cleaning_removes_only_exact_duplicate_lines(res):
    q = res["quality"]
    assert q["rows_clean"] == q["rows_raw"] - q["duplicate_lines_removed"]
    assert q["sales_clean"] == pytest.approx(q["sales_raw"] - q["duplicate_sales_removed"], abs=0.01)


def test_delivered_columns_are_arithmetically_consistent(res):
    q = res["quality"]
    assert q["line_sales_formula_violations"] == 0
    assert q["line_cost_formula_violations"] == 0
    assert q["unit_price_varies_within_item"] == 0


def test_bulk_threshold_sits_in_a_real_gap_of_the_quantity_distribution(res):
    q = res["quality"]
    assert q["max_quantity_below_bulk_threshold"] * 5 <= q["min_quantity_at_or_above_bulk_threshold"]


def test_store_figures_add_up_to_chain(res):
    s = res["q1"]["stores"]
    assert s.Bills.sum() == res["q2"]["bills"]
    assert s.Sales.sum() == pytest.approx(res["quality"]["sales_clean"], abs=0.01)
    c = res["q3"]["categories"]
    assert c.Sales_Q1.sum() + c.Sales_Q4.sum() <= s.Sales.sum()


def test_focus_store_is_the_one_with_largest_profit_upside(res):
    s = res["q1"]["stores"]
    assert s.ProfitUpside.idxmax() == res["q1"]["focus"]


def test_mix_and_rate_effects_add_up_to_the_profit_upside(res):
    q1 = res["q1"]
    upside = res["q1"]["stores"].loc[q1["focus"], "ProfitUpside"]
    assert q1["focus_extra_profit_same_sales"] + q1["focus_rate_effect"] == pytest.approx(upside, abs=0.01)


def test_q2_gap_explanations_shrink_the_gap(res):
    steps = res["q2"]["gap_steps"]
    assert steps.Gap.iloc[0] == pytest.approx(res["q2"]["gap"])
    assert steps.Gap.iloc[-1] < steps.Gap.iloc[0]
    assert (steps.ShareOfGapRemoved >= -1e-9).all() and (steps.ShareOfGapRemoved <= 1).all()


def test_q5_second_kpi_counts_exactly_the_below_cost_sale_lines(res):
    df = res["df"]
    n = ((df.SaleQuantity > 0) & (df.LineSales < df.LineCost)).sum()
    kpi = res["q5"][res["q5"].KPI.str.startswith("2.")].Baseline.iloc[0]
    assert kpi == pytest.approx(n / 12)
    q = res["quality"]   # the two data-quality flags are the two ways a line ends up below cost
    assert n == q["discount_ge_gross_sale_lines"] + q["unit_cost_above_unit_price"]


def test_category_contributions_add_up_to_the_mix_gap(res):
    q1 = res["q1"]
    assert q1["categories"].Contribution.sum() == pytest.approx(q1["focus_extra_profit_same_sales"], abs=0.01)
    assert q1["categories"].FocusMix.sum() == pytest.approx(1)
    assert q1["categories"].RestMix.sum() == pytest.approx(1)


def test_bill_numbers_do_not_overlap_between_stores_in_this_file():
    # True here, but the code still keys receipts on StoreCode + BillNum, as the brief requires.
    raw = load_raw()
    assert raw.groupby("BillNum")["StoreCode"].nunique().max() == 1


def test_sql_q1_matches_pandas(res, sql):
    py = res["q1"]["stores"].sort_index()
    q = sql["q1_stores"].set_index("StoreName").sort_index()
    for col in ["Sales", "Profit", "Bills"]:
        assert q[col].to_numpy() == pytest.approx(py[col].to_numpy(), abs=0.01)
    assert q["Margin"].to_numpy() == pytest.approx(py["Margin"].to_numpy(), abs=1e-5)
    assert q["AvgBill"].to_numpy() == pytest.approx(py["AvgBill"].to_numpy(), abs=1e-3)


def test_sql_q2_matches_pandas(res, sql):
    row = sql["q2_receipt_mean_median"].iloc[0]
    assert row.Bills == res["q2"]["bills"]
    assert row.Mean == pytest.approx(res["q2"]["mean"], abs=1e-3)
    assert row.Median == pytest.approx(res["q2"]["median"], abs=1e-3)


def test_sql_q3_matches_pandas(res, sql):
    py = res["q3"]["categories"].sort_index()
    q = sql["q3_category_q1_vs_q4"].set_index("Category").sort_index()
    for col in ["Sales_Q1", "Sales_Q4", "Profit_Q1", "Profit_Q4"]:
        assert q[col].to_numpy() == pytest.approx(py[col].to_numpy(), abs=0.01)
    assert q["SalesChangePct"].to_numpy() == pytest.approx(py["SalesChangePct"].to_numpy(), abs=1e-5)


def test_sql_q4_matches_pandas(res, sql):
    py = res["q4"]["attach_rates"]
    long = sql["q4_attach_rates"].pivot(index="StoreName", columns="Category", values="Attach").fillna(0)
    cats = list(long.columns)
    assert set(cats) <= set(py.columns)
    assert long.sort_index().to_numpy() == pytest.approx(py.sort_index()[cats].to_numpy(), abs=1e-5)


@pytest.mark.skipif(not config.EXCEL_PATH.exists(), reason="workbook not built")
def test_excel_formulas_match_python(res):
    """Every row of the workbook's Checks sheet compares an Excel formula result with the Python result."""
    from openpyxl import load_workbook
    wb = load_workbook(config.EXCEL_PATH, data_only=True)
    checks = wb["Checks"]
    diffs = [(checks.cell(r, 1).value, checks.cell(r, 4).value) for r in range(5, 40)
             if isinstance(checks.cell(r, 4).value, (int, float))]
    if not diffs:
        pytest.skip("workbook has no cached values - run excel/recalc.ps1")
    assert len(diffs) >= 10
    bad = [(n, d) for n, d in diffs if abs(d) > 1e-4]
    assert not bad, f"Excel and Python disagree: {bad}"
    assert wb["Assumptions"]["B6"].value == res["q1"]["focus"], "Excel and Python chose different focus stores"


# ============================================================== ADAPTIVITY
def test_conclusions_follow_the_numbers(res):
    """Make another store the weak one (its costs +40%) and the whole analysis must re-point at it."""
    df = res["df"].copy()
    victim = "Ashdod"
    assert res["q1"]["focus"] != victim
    m = df.StoreName == victim
    df.loc[m, "LineCost"] = df.loc[m, "LineCost"] * 1.4
    df["Profit"] = df["LineSales"] - df["LineCost"]

    q1 = questions.q1_stores(df)
    assert q1["focus"] == victim
    assert q1["driver_categories"] == []                      # a cost problem is not a category-mix problem ...
    assert q1["focus_rate_effect"] > 2 * abs(q1["focus_extra_profit_same_sales"])   # ... and the split says so
    q4 = questions.q4_open_investigation(df)
    q5 = questions.q5_kpi_baselines(df)
    assert q4["subject"] == victim                            # Q4 and Q5 find the same store on their own
    assert any(victim in k for k in q5.KPI)
    assert victim in q4["common_customer_category_mix"].index
    assert all(victim in tbl.index for tbl in q4["share_by_quarter_of_chain_movers"].values())
    # the cost change affects margins, not the quantities, so the receipt statistics are untouched
    assert questions.q2_receipt_distribution(df)["mean"] == pytest.approx(res["q2"]["mean"])


def test_each_question_is_a_function_of_the_data_only(res):
    """The brief's five questions are independent: no question receives another question's result."""
    import inspect
    for fn in (questions.q1_stores, questions.q2_receipt_distribution, questions.q3_category_change,
               questions.q4_open_investigation, questions.q5_kpi_baselines):
        assert list(inspect.signature(fn).parameters) == ["df"], fn.__name__
    # ... and therefore the answer to Q3 does not change when the store comparison does
    df = res["df"].copy()
    df.loc[df.StoreName == "Ashdod", "LineCost"] *= 1.4
    df["Profit"] = df["LineSales"] - df["LineCost"]
    same = questions.q3_category_change(df)["categories"][["Sales_Q1", "Sales_Q4"]]
    assert same.to_numpy() == pytest.approx(res["q3"]["categories"][["Sales_Q1", "Sales_Q4"]].to_numpy())


def test_shifting_sales_between_categories_changes_the_biggest_mover(res):
    df = res["df"].copy()
    assert res["q3"]["biggest_growth"] != "Toys"
    toys_q4 = (df.Category == "Toys") & (df.Quarter == 4)
    df.loc[toys_q4, ["LineSales", "LineCost", "SaleQuantity", "Profit"]] *= 20  # an obvious new winner
    q3 = questions.q3_category_change(df)
    assert q3["biggest_growth"] == "Toys"


# ============================================================== REFERENCE FIGURES (delivered dataset)
def test_reference_chain_totals(res):
    s = res["q1"]["stores"]
    pin("chain sales", s.Sales.sum(), 3_520_099.61, 0.01)
    pin("chain profit", s.Profit.sum(), 1_157_552.43, 0.01)
    pin("receipts", s.Bills.sum(), 15_588, 0)
    pin("duplicate lines removed", res["quality"]["duplicate_lines_removed"], 59, 0)


def test_reference_q1_focus_store_and_mix_effect(res):
    q1 = res["q1"]
    assert q1["focus"] == "Haifa", f"focus store is now {q1['focus']!r}; the delivered answers.md is about Haifa"
    pin("focus margin", q1["focus_margin_actual"], 0.2792, 1e-3)
    pin("focus margin at rest-of-chain mix", q1["focus_margin_at_rest_mix"], 0.3425, 1e-3)
    pin("extra profit", q1["focus_extra_profit_same_sales"], 52_945, 1)
    assert q1["driver_categories"] == ["Wet Food", "Snacks"]
    assert q1["overweighted_category"] == "Dry Food"
    cat = q1["categories"]  # per-category margins are the same in the focus store and elsewhere
    assert (cat.FocusMargin - cat.RestMargin).abs().max() < 0.01


def test_reference_q2_mean_median(res):
    q2 = res["q2"]
    pin("mean receipt", q2["mean"], 225.82, 0.01)
    pin("median receipt", q2["median"], 186.545, 0.001)
    pin("mean without bulk receipts", q2["mean_ex_bulk"], 214.74, 0.01)
    pin("median without bulk receipts", q2["median_ex_bulk"], 185.73, 0.01)
    pin("share of the gap explained by bulk receipts", q2["gap_explained_by_bulk"], 0.26, 0.02)


def test_reference_q1_store_trends(res):
    s = res["q1"]["stores"]
    pin("Ashdod sales Q4 vs Q1", s.loc["Ashdod", "SalesChangeQ4vsQ1"], -0.072, 1e-3)
    pin("Haifa sales Q4 vs Q1", s.loc["Haifa", "SalesChangeQ4vsQ1"], -0.018, 1e-3)
    pin("Kiryat Ono sales Q4 vs Q1", s.loc["Kiryat Ono", "SalesChangeQ4vsQ1"], 0.075, 1e-3)
    pin("Haifa profit upside", s.loc["Haifa", "ProfitUpside"], 54_423, 1)


def test_reference_q2_gap_explanations(res):
    q2 = res["q2"]
    assert q2["priciest_category"] == "Dry Food"
    steps = q2["gap_steps"]
    pin("gap within receipts of the same number of lines", steps.Gap.iloc[2], 32.02, 0.01)
    pin("gap within comparable receipts", steps.Gap.iloc[3], 13.22, 0.01)
    pin("share of the gap removed", steps.ShareOfGapRemoved.iloc[3], 0.663, 1e-3)


def test_reference_q5_baselines(res):
    kpi = res["q5"].set_index("KPI").Baseline
    pin("sale lines below cost per month", kpi["2. Sale lines below cost per month"], 19.5, 1e-9)
    pin("Haifa wet-food attach", kpi["1. Wet Food attach rate, Haifa"], 0.391, 1e-3)


def test_reference_q3_mix_shift(res):
    q3, c = res["q3"], res["q3"]["categories"]
    assert (q3["biggest_decline"], q3["biggest_growth"]) == ("Dry Food", "Wet Food")
    pin("dry food change", c.loc["Dry Food", "SalesChangePct"], -0.276, 1e-3)
    pin("wet food change", c.loc["Wet Food", "SalesChangePct"], 1.330, 1e-3)
    pin("chain sales change", q3["total_sales_change_pct"], 0.012, 1e-3)
    pin("chain profit change", q3["total_profit_change_pct"], 0.125, 1e-3)
    assert (c.Margin_Q1 - c.Margin_Q4).abs().max() < 0.01  # category margins are stable


def test_reference_q4_same_customers_buy_differently_in_focus_store(res):
    q4 = res["q4"]
    mix = q4["common_customer_category_mix"]["Dry Food"]
    pin("dry-food share of common customers in focus store", mix["Haifa"], 0.794, 1e-3)
    pin("dry-food share of common customers elsewhere", mix["Rest of chain"], 0.497, 1e-3)
    pin("common customers", q4["common_customers"], 1294, 0)
    pin("focus wet-food attach", q4["attach_rates"].loc["Haifa", "Wet Food"], 0.391, 1e-3)
    pin("rest wet-food attach", q4["rest_attach_rates"]["Wet Food"], 0.665, 1e-3)
    pin("rest snacks attach", q4["rest_attach_rates"]["Snacks"], 0.465, 1e-3)
    assert q4["list_price_varies_across_stores_items"] == 0
